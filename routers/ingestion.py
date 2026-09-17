"""
FEATURE 1 — Intelligent Document Ingestion Engine
=================================================
Mounted at /api/ingest

Pipeline: OCR -> LLM entity extraction -> SQLite -> ChromaDB -> plant graph -> background monitor.
Zero external network calls. OCR is local (EasyOCR), extraction is local (Ollama).
"""

from __future__ import annotations

import io
import json
import os
import re
import time
from datetime import datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, BackgroundTasks, File, Form, HTTPException, Query, UploadFile

import llm
from data_layer import (
    chunk_text, get_collection, get_db, get_plant_graph, log_audit,
    parse_tags, rows_to_dicts, save_plant_graph, upsert_equipment_node,
)

router = APIRouter(prefix="/api/ingest", tags=["ingestion"])

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
TEXT_EXT = {".txt", ".md", ".csv", ".log"}

# Equipment tag pattern used as a safety net alongside LLM extraction:
# FT-201, V-301, HE-301A, P-101B, 10-PSV-4021
TAG_RE = re.compile(r"\b(?:\d{1,2}-)?([A-Z]{1,4})-?(\d{2,4}[A-Z]?)\b")
KNOWN_PREFIXES = {
    "V", "D", "T", "C", "HE", "E", "P", "K", "F", "H", "R",
    "FT", "PT", "TT", "LT", "PI", "TI", "FI", "LI", "PSV", "CV", "TK",
}
STANDARD_RE = re.compile(
    r"\b(OISD[- ]?(?:STD[- ]?)?\d{2,3}|IS[- ]?\d{3,5}(?:[-:]\d{4})?|"
    r"ASME\s+(?:SEC(?:TION)?\s+)?[IVXL]+(?:\s+DIV\s*\d)?|API\s?\d{3,4}|"
    r"IEC\s?\d{4,5}|ISO\s?\d{3,5}|BS\s?\d{3,4})\b",
    re.IGNORECASE,
)

# --------------------------------------------------------------------------------------
# Lazily loaded OCR reader (EasyOCR model load is ~10s, so do it once, on demand)
# --------------------------------------------------------------------------------------

_ocr_reader = None


def get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        import easyocr
        _ocr_reader = easyocr.Reader(["en", "hi"], gpu=False)
    return _ocr_reader


# --------------------------------------------------------------------------------------
# STEP 1 — OCR & text extraction
# --------------------------------------------------------------------------------------

def extract_from_pdf(data: bytes) -> Dict[str, Any]:
    """pdfplumber first; fallback to pypdf; fallback to pypdfium2/pdf2image + EasyOCR for scans."""
    pages: List[str] = []
    method = "pdfplumber"
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
                for table in (page.extract_tables() or []):
                    for row in table:
                        cells = [str(c).strip() for c in row if c]
                        if cells:
                            pages.append(" | ".join(cells))
    except Exception as exc:
        print(f"[ingest] pdfplumber failed: {exc}")

    text = "\n".join(p for p in pages if p).strip()

    if not text:
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            pypdf_pages = [page.extract_text() or "" for page in reader.pages]
            pypdf_text = "\n".join(pypdf_pages).strip()
            if pypdf_text:
                text = pypdf_text
                method = "pypdf"
                if not pages:
                    pages = pypdf_pages
        except Exception as exc:
            print(f"[ingest] pypdf fallback failed: {exc}")

    if len(text.strip()) < 30:
        # Scanned document -> rasterise and OCR every page via pypdfium2 (no poppler needed)
        try:
            import pypdfium2 as pdfium
            import numpy as np
            reader = get_ocr_reader()
            ocr_pages = []
            doc = pdfium.PdfDocument(io.BytesIO(data))
            for page in doc:
                pil_img = page.render(scale=2.0).to_pil().convert("RGB")
                result = reader.readtext(np.array(pil_img), detail=0, paragraph=True)
                ocr_pages.append("\n".join(result))
            ocr_text = "\n".join(ocr_pages).strip()
            if len(ocr_text) > len(text):
                text, method = ocr_text, "pypdfium2+easyocr"
                if not pages:
                    pages = ocr_pages
        except Exception as exc:
            print(f"[ingest] pypdfium2 OCR fallback failed: {exc}")
            try:
                from pdf2image import convert_from_bytes
                import numpy as np
                reader = get_ocr_reader()
                ocr_pages = []
                for image in convert_from_bytes(data, dpi=200):
                    result = reader.readtext(np.array(image), detail=0, paragraph=True)
                    ocr_pages.append("\n".join(result))
                ocr_text = "\n".join(ocr_pages).strip()
                if len(ocr_text) > len(text):
                    text, method = ocr_text, "pdf2image+easyocr"
            except Exception as exc2:
                print(f"[ingest] pdf2image OCR fallback failed: {exc2}")

    return {"text": text, "method": method, "page_count": len(pages)}


def extract_from_image(data: bytes) -> Dict[str, Any]:
    import numpy as np
    from PIL import Image

    image = Image.open(io.BytesIO(data)).convert("RGB")
    width, height = image.size
    # Wide or tall canvas -> most likely a P&ID / isometric drawing sheet.
    ratio = max(width / max(height, 1), height / max(width, 1))
    reader = get_ocr_reader()
    lines = reader.readtext(np.array(image), detail=0, paragraph=True)
    return {
        "text": "\n".join(lines).strip(),
        "method": "easyocr",
        "engineering_drawing": ratio > 2.0,
        "dimensions": f"{width}x{height}",
    }


def extract_text(filename: str, data: bytes) -> Dict[str, Any]:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_from_pdf(data)
    if ext in IMAGE_EXT:
        return extract_from_image(data)
    if ext in TEXT_EXT:
        return {"text": data.decode("utf-8", errors="replace"), "method": "plaintext"}
    if ext == ".docx":
        try:
            import docx
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return {"text": "\n".join(parts).strip(), "method": "python-docx"}
        except Exception as exc:
            raise HTTPException(400, f"Could not read .docx: {exc}")
    # Last resort: treat as text.
    return {"text": data.decode("utf-8", errors="replace"), "method": "raw_decode"}


# --------------------------------------------------------------------------------------
# STEP 2 — Entity extraction via local Ollama
# --------------------------------------------------------------------------------------

EXTRACTION_SYSTEM = """You are an industrial document parser for Indian refineries and PSUs.
Extract ALL of the following from the text and return ONLY valid JSON, nothing else:
{
  "equipment_tags": ["list of tags like FT-201, V-301, P-101, HE-201"],
  "measurements": [{"tag": "", "parameter": "", "value": "", "unit": "", "date": ""}],
  "standard_references": ["OISD-118", "IS 2825", "ASME VIII"],
  "document_type": "inspection_report|approval_note|sop|drawing|correspondence|calculation",
  "document_date": "YYYY-MM-DD or null",
  "key_findings": ["list of factual findings only"],
  "action_items": ["list of action items if any"],
  "severity_indicators": ["words like critical, urgent, overdue, failed, exceeded"]
}
Extract only what is explicitly stated. Return empty arrays for missing fields."""

EMPTY_EXTRACTION = {
    "equipment_tags": [], "measurements": [], "standard_references": [],
    "document_type": None, "document_date": None, "key_findings": [],
    "action_items": [], "severity_indicators": [],
}


def regex_fallback(text: str) -> Dict[str, List[str]]:
    """Deterministic backstop so ingestion never returns nothing when the model is down."""
    tags = set()
    for match in TAG_RE.finditer(text.upper()):
        prefix, number = match.group(1), match.group(2)
        if prefix in KNOWN_PREFIXES:
            tags.add(f"{prefix}-{number}")
    standards = {re.sub(r"\s+", "-", m.group(0).upper().replace("STD-", ""))
                 for m in STANDARD_RE.finditer(text)}
    return {"equipment_tags": sorted(tags), "standard_references": sorted(standards)}


# Deterministic measurement patterns: "Wall thickness: 12.8 mm", "Vibration reading = 4.2 mm/s"
MEASUREMENT_RE = re.compile(
    r"([A-Za-z][A-Za-z \-/()]{2,40}?)\s*[:=]\s*(-?\d+(?:\.\d+)?)\s*"
    r"(mm/s|mm/yr|mm|barg|bar|kg/cm2|kPa|MPa|psi|deg\s?C|°C|%|m3/hr|m³/hr|kL/hr|"
    r"LPM|RPM|micron|µm|ppm|Hz|kW|V|A)?",
    re.IGNORECASE,
)
DATE_LINE_RE = re.compile(
    r"\b(\d{4}-\d{2}-\d{2})\b|\b(\d{1,2}[-/]\d{1,2}[-/]\d{4})\b"
)
NOISE_PARAMS = {
    "date", "status", "note", "notes", "report", "tag", "equipment", "equipment_tag",
    "inspected_by", "reviewed_by", "logged_by", "location", "unit", "prepared_by",
    "approved_by", "document", "reference", "page", "revision", "sheet", "title",
    "service", "remarks", "observations", "purpose", "scope", "shell_material",
}
# Bibliographic / administrative fields are never engineering measurements.
NOISE_TOKENS = ("_no", "_due", "number", "reference", "revision", "version", "_date",
                "_by", "_code", "_method", "material", "standard", "interval")


def is_measurement_param(parameter: str, value: float, unit: str) -> bool:
    """Filter out document metadata that merely looks like 'label: number'."""
    if not parameter or len(parameter) < 3 or parameter in NOISE_PARAMS:
        return False
    if any(token in parameter for token in NOISE_TOKENS):
        return False
    # A bare four-digit year with no unit is a date fragment, not a reading.
    if not unit and float(value).is_integer() and 1900 <= value <= 2100:
        return False
    return True


# "Equipment Tag: HE-301" style declarations define whose readings follow.
TAG_DECLARATION_RE = re.compile(
    r"^\s*(?:equipment\s*tag|tag\s*no\.?|tag|equipment)\s*[:=]\s*([A-Z]{1,4}-?\d{2,4}[A-Z]?)",
    re.IGNORECASE,
)


def find_subject_tag(text: str) -> Optional[str]:
    """The tag this document is *about*, taken from its header declaration."""
    for line in text.splitlines()[:40]:
        match = TAG_DECLARATION_RE.match(line)
        if match:
            tag = match.group(1).upper()
            if "-" not in tag:
                tag = re.sub(r"^([A-Z]+)(\d+)$", r"\1-\2", tag)
            return tag
    return None


def find_document_date(text: str) -> Optional[str]:
    """First explicit date in the header block, normalised to YYYY-MM-DD."""
    for line in text.splitlines()[:40]:
        match = DATE_LINE_RE.search(line)
        if not match:
            continue
        if match.group(1):
            return match.group(1)
        parts = re.split(r"[-/]", match.group(2))
        if len(parts) == 3:
            day, month, year = parts
            try:
                return f"{year}-{int(month):02d}-{int(day):02d}"
            except ValueError:
                return None
    return None


def regex_measurements(text: str, tags: List[str],
                       primary_tag: Optional[str] = None) -> List[Dict[str, Any]]:
    """Parse 'parameter: value unit' lines so measurements survive a model outage."""
    default_tag = primary_tag or find_subject_tag(text) or (tags[0] if tags else "UNTAGGED")
    doc_date = find_document_date(text)

    found: List[Dict[str, Any]] = []
    current_tag = default_tag
    for line in text.splitlines():
        # Only an explicit declaration reassigns ownership. A passing mention of another
        # tag (e.g. "Service: crude transfer to HE-301") must not steal the readings.
        declared = TAG_DECLARATION_RE.match(line)
        if declared:
            current_tag = declared.group(1).upper()
        match = MEASUREMENT_RE.search(line)
        if not match:
            continue
        parameter = normalise_parameter(match.group(1))
        value = float(match.group(2))
        unit = (match.group(3) or "").strip()
        if not is_measurement_param(parameter, value, unit):
            continue
        found.append({"tag": current_tag, "parameter": parameter, "value": value,
                      "unit": unit, "date": doc_date})
    return found


def normalise_parameter(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (name or "").strip().lower()).strip("_")


def coerce_float(value: Any) -> Optional[float]:
    if isinstance(value, (int, float)):
        return float(value)
    if not value:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group(0)) if match else None


def extract_entities(text: str, feature: str = "ingestion",
                     primary_tag: Optional[str] = None) -> Dict[str, Any]:
    """LLM extraction with a regex safety net merged in."""
    snippet = text[:12000]
    primary_tag = (primary_tag or "").strip().upper() or find_subject_tag(text)
    result = dict(EMPTY_EXTRACTION)
    try:
        parsed = llm.generate_json(
            f"TEXT TO PARSE:\n---\n{snippet}\n---",
            system=EXTRACTION_SYSTEM, feature=feature, task_type="extraction",
            default=None, retries=1, num_predict=2200,
        )
        if isinstance(parsed, dict):
            result.update({k: parsed.get(k, v) for k, v in EMPTY_EXTRACTION.items()})
    except llm.ModelUnavailable as exc:
        print(f"[ingest] model unavailable, regex-only extraction: {exc}")
    except Exception as exc:
        print(f"[ingest] LLM entity extraction failed, falling back to regex: {exc}")

    fallback = regex_fallback(text)
    tags = {str(t).strip().upper() for t in (result.get("equipment_tags") or []) if t}
    tags.update(fallback["equipment_tags"])
    if primary_tag:
        tags.add(primary_tag)
    # Subject tag first: downstream code treats index 0 as the primary equipment.
    ordered = sorted(tags)
    if primary_tag and primary_tag in ordered:
        ordered.remove(primary_tag)
        ordered.insert(0, primary_tag)
    result["equipment_tags"] = ordered
    result["primary_tag"] = primary_tag or (ordered[0] if ordered else None)
    if not result.get("document_date"):
        result["document_date"] = find_document_date(text)

    standards = {str(s).strip().upper() for s in (result.get("standard_references") or []) if s}
    standards.update(fallback["standard_references"])
    result["standard_references"] = sorted(standards)

    # Normalise measurements
    clean: List[Dict[str, Any]] = []
    for m in (result.get("measurements") or []):
        if not isinstance(m, dict):
            continue
        value = coerce_float(m.get("value"))
        if value is None:
            continue
        clean.append({
            "tag": str(m.get("tag") or "").strip().upper()
                   or result["primary_tag"] or "UNTAGGED",
            "parameter": normalise_parameter(m.get("parameter")),
            "value": value,
            "unit": (m.get("unit") or "").strip(),
            "date": (m.get("date") or result.get("document_date") or "") or None,
        })
    clean = [m for m in clean if is_measurement_param(m["parameter"], m["value"], m["unit"])]

    # Merge in deterministic parses for anything the model missed.
    existing = {(m["tag"], m["parameter"]) for m in clean}
    for m in regex_measurements(text, result["equipment_tags"], result["primary_tag"]):
        if (m["tag"], m["parameter"]) not in existing:
            clean.append(m)
            existing.add((m["tag"], m["parameter"]))

    result["measurements"] = clean
    return result


# --------------------------------------------------------------------------------------
# Endpoints
# --------------------------------------------------------------------------------------

@router.post("/upload")
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    doc_type: Optional[str] = Form(None),
    equipment_tag: Optional[str] = Form(None),
    uploaded_by: Optional[str] = Form("engineer"),
):
    """Full six-step ingestion pipeline for one document."""
    started = time.perf_counter()
    data = await file.read()
    if not data:
        raise HTTPException(400, "Empty file received.")

    safe_name = os.path.basename(file.filename or "upload.bin")
    stored_path = os.path.join(UPLOAD_DIR, f"{int(time.time())}_{safe_name}")
    with open(stored_path, "wb") as fh:
        fh.write(data)

    # ---------- STEP 1: OCR / text extraction ----------
    extracted = extract_text(safe_name, data)
    raw_text = extracted.get("text") or ""
    if not raw_text.strip():
        raise HTTPException(
            422,
            "No readable text or characters could be detected in this document. "
            "If this is a scanned document or drawing, ensure the file is clear, readable, and not password-protected.",
        )

    # ---------- STEP 2: entity extraction ----------
    entities = extract_entities(raw_text, primary_tag=equipment_tag)
    if equipment_tag:
        manual = equipment_tag.strip().upper()
        if manual not in entities["equipment_tags"]:
            entities["equipment_tags"].insert(0, manual)

    resolved_type = doc_type or entities.get("document_type") or "correspondence"
    tags_csv = ",".join(entities["equipment_tags"])
    standards_csv = ",".join(entities["standard_references"])
    doc_date = entities.get("document_date") or None

    # ---------- STEP 3: SQLite ----------
    conn = get_db()
    try:
        cur = conn.execute(
            """INSERT INTO documents
               (filename, upload_time, doc_type, equipment_tags, standard_refs, raw_text,
                processed, document_date, uploaded_by, key_findings, action_items, engineering_drawing)
               VALUES (?, ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, ?)""",
            (safe_name, datetime.now().isoformat(), resolved_type, tags_csv, standards_csv,
             raw_text, doc_date, uploaded_by,
             json.dumps(entities.get("key_findings") or []),
             json.dumps(entities.get("action_items") or []),
             1 if extracted.get("engineering_drawing") else 0),
        )
        doc_id = cur.lastrowid

        for m in entities["measurements"]:
            conn.execute(
                """INSERT INTO measurements
                   (doc_id, equipment_tag, parameter, value, unit, measurement_date, source_line)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (doc_id, m["tag"], m["parameter"], m["value"], m["unit"],
                 m["date"] or doc_date or datetime.now().date().isoformat(),
                 f"{m['parameter']} = {m['value']} {m['unit']}".strip()),
            )

        entity_rows = 0
        for etype, values in (
            ("equipment_tag", entities["equipment_tags"]),
            ("standard_reference", entities["standard_references"]),
            ("key_finding", entities.get("key_findings") or []),
            ("action_item", entities.get("action_items") or []),
            ("severity_indicator", entities.get("severity_indicators") or []),
        ):
            for idx, value in enumerate(values):
                if not value:
                    continue
                conn.execute(
                    """INSERT INTO entities (doc_id, entity_type, entity_value, context, position)
                       VALUES (?, ?, ?, ?, ?)""",
                    (doc_id, etype, str(value)[:1000], resolved_type, idx),
                )
                entity_rows += 1
        conn.commit()
    finally:
        conn.close()

    # ---------- STEP 4: ChromaDB ----------
    chunks = chunk_text(raw_text, 512, 128)
    if chunks:
        try:
            get_collection("documents").upsert(
                ids=[f"doc{doc_id}_c{i}" for i in range(len(chunks))],
                documents=chunks,
                metadatas=[{
                    "doc_id": doc_id, "filename": safe_name, "doc_type": resolved_type,
                    "equipment_tags": tags_csv, "document_date": doc_date or "",
                    "chunk_index": i,
                } for i in range(len(chunks))],
            )
        except Exception as exc:
            print(f"[ingest] vector upsert failed: {exc}")

    # Structured measurements also go into the history collection for pattern matching.
    if entities["measurements"]:
        try:
            get_collection("measurements_history").upsert(
                ids=[f"doc{doc_id}_m{i}" for i in range(len(entities["measurements"]))],
                documents=[
                    f"{m['tag']} {m['parameter'].replace('_', ' ')} measured {m['value']} {m['unit']} "
                    f"on {m['date'] or doc_date or 'unknown date'} in {resolved_type} {safe_name}"
                    for m in entities["measurements"]
                ],
                metadatas=[{
                    "doc_id": doc_id, "equipment_tag": m["tag"], "parameter": m["parameter"],
                    "value": m["value"], "unit": m["unit"], "date": m["date"] or doc_date or "",
                } for m in entities["measurements"]],
            )
        except Exception as exc:
            print(f"[ingest] measurement vector upsert failed: {exc}")

    # If it is an SOP, mirror it into the SOP library for the compliance engine.
    if resolved_type == "sop":
        conn = get_db()
        try:
            conn.execute(
                """INSERT INTO sop_documents
                   (filename, sop_code, standard_refs, version, effective_date, raw_text, active, domain)
                   VALUES (?, ?, ?, ?, ?, ?, 1, ?)""",
                (safe_name, (entities["equipment_tags"] or [safe_name])[0], standards_csv,
                 "1.0", doc_date, raw_text, resolved_type),
            )
            conn.commit()
        finally:
            conn.close()

    # ---------- STEP 5: plant graph ----------
    G = get_plant_graph()
    created = 0
    for tag in entities["equipment_tags"]:
        if upsert_equipment_node(G, tag, last_inspection_date=doc_date):
            created += 1
    # Tags co-occurring in one document are process-related; record a weak association edge.
    tags = entities["equipment_tags"]
    if len(tags) > 1:
        primary = tags[0]
        for other in tags[1:]:
            if not G.has_edge(primary, other):
                G.add_edge(primary, other, connection_type="co_referenced",
                           line_number=None, source_doc=doc_id, inferred=True)
    save_plant_graph(G)

    elapsed = int((time.perf_counter() - started) * 1000)

    log_audit("document_uploaded", uploaded_by, doc_id, "ingestion", {
        "filename": safe_name, "doc_type": resolved_type,
        "extraction_method": extracted.get("method"),
        "equipment_tags": entities["equipment_tags"],
        "measurements_found": len(entities["measurements"]),
        "entities_found": entity_rows, "processing_time_ms": elapsed,
    })

    # ---------- STEP 6: background passive analysis ----------
    from routers.monitoring import analyze_document_task
    background_tasks.add_task(analyze_document_task, doc_id)

    return {
        "doc_id": doc_id,
        "filename": safe_name,
        "doc_type": resolved_type,
        "extraction_method": extracted.get("method"),
        "engineering_drawing": bool(extracted.get("engineering_drawing")),
        "equipment_tags": entities["equipment_tags"],
        "standard_references": entities["standard_references"],
        "measurements_found": len(entities["measurements"]),
        "measurements": entities["measurements"],
        "entities_found": entity_rows,
        "key_findings": entities.get("key_findings") or [],
        "action_items": entities.get("action_items") or [],
        "graph_nodes_updated": len(entities["equipment_tags"]),
        "graph_nodes_created": created,
        "chunks_indexed": len(chunks),
        "text_length": len(raw_text),
        "processing_time_ms": elapsed,
        "monitoring": "queued",
    }


@router.get("/documents")
async def list_documents(
    equipment_tag: Optional[str] = Query(None),
    doc_type: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    limit: int = Query(200, le=1000),
):
    """Document register with metadata (raw text excluded for payload size)."""
    sql = """
        SELECT d.id, d.filename, d.upload_time, d.doc_type, d.equipment_tags, d.standard_refs,
               d.document_date, d.uploaded_by, d.processed, d.engineering_drawing, d.superseded,
               d.key_findings, d.action_items,
               (SELECT COUNT(*) FROM measurements m WHERE m.doc_id = d.id) AS measurements_found,
               (SELECT COUNT(*) FROM entities e WHERE e.doc_id = d.id)     AS entities_found,
               LENGTH(d.raw_text) AS text_length
        FROM documents d WHERE 1=1
    """
    params: List[Any] = []
    if equipment_tag:
        sql += " AND UPPER(d.equipment_tags) LIKE ?"
        params.append(f"%{equipment_tag.upper()}%")
    if doc_type:
        sql += " AND d.doc_type = ?"
        params.append(doc_type)
    if date_from:
        sql += " AND COALESCE(d.document_date, d.upload_time) >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND COALESCE(d.document_date, d.upload_time) <= ?"
        params.append(date_to)
    sql += " ORDER BY COALESCE(d.document_date, d.upload_time) DESC LIMIT ?"
    params.append(limit)

    conn = get_db()
    try:
        docs = rows_to_dicts(conn.execute(sql, params).fetchall())
        for d in docs:
            d["equipment_tags"] = parse_tags(d.get("equipment_tags"))
            d["standard_refs"] = parse_tags(d.get("standard_refs"))
            for key in ("key_findings", "action_items"):
                try:
                    d[key] = json.loads(d.get(key) or "[]")
                except json.JSONDecodeError:
                    d[key] = []
            d["alerts_triggered"] = conn.execute(
                "SELECT COUNT(*) c FROM alerts WHERE evidence LIKE ?", (f'%"doc_id": {d["id"]}%',)
            ).fetchone()["c"]
        return {"total": len(docs), "documents": docs}
    finally:
        conn.close()


@router.get("/document/{doc_id}")
async def get_document(doc_id: int):
    """Full record: raw text, measurements, entities, alerts."""
    conn = get_db()
    try:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
        if not row:
            raise HTTPException(404, f"Document {doc_id} not found.")
        doc = dict(row)
        doc["equipment_tags"] = parse_tags(doc.get("equipment_tags"))
        doc["standard_refs"] = parse_tags(doc.get("standard_refs"))
        for key in ("key_findings", "action_items"):
            try:
                doc[key] = json.loads(doc.get(key) or "[]")
            except json.JSONDecodeError:
                doc[key] = []
        doc["measurements"] = rows_to_dicts(conn.execute(
            "SELECT * FROM measurements WHERE doc_id = ? ORDER BY measurement_date", (doc_id,)
        ).fetchall())
        doc["entities"] = rows_to_dicts(conn.execute(
            "SELECT * FROM entities WHERE doc_id = ? ORDER BY entity_type, position", (doc_id,)
        ).fetchall())
        return doc
    finally:
        conn.close()


@router.delete("/document/{doc_id}")
async def delete_document(doc_id: int, deleted_by: str = Query("engineer")):
    conn = get_db()
    try:
        if not conn.execute("SELECT 1 FROM documents WHERE id = ?", (doc_id,)).fetchone():
            raise HTTPException(404, f"Document {doc_id} not found.")
        conn.execute("DELETE FROM measurements WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM entities WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM claims WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        conn.commit()
    finally:
        conn.close()
    log_audit("document_deleted", deleted_by, doc_id, "ingestion", {"doc_id": doc_id})
    return {"deleted": doc_id}


@router.get("/stats")
async def ingestion_stats():
    conn = get_db()
    try:
        return {
            "documents": conn.execute("SELECT COUNT(*) c FROM documents").fetchone()["c"],
            "measurements": conn.execute("SELECT COUNT(*) c FROM measurements").fetchone()["c"],
            "entities": conn.execute("SELECT COUNT(*) c FROM entities").fetchone()["c"],
            "equipment_tracked": conn.execute(
                "SELECT COUNT(DISTINCT equipment_tag) c FROM measurements").fetchone()["c"],
            "by_type": {r["doc_type"] or "unknown": r["c"] for r in conn.execute(
                "SELECT doc_type, COUNT(*) c FROM documents GROUP BY doc_type").fetchall()},
        }
    finally:
        conn.close()
